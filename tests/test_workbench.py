import asyncio
import io
import json
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from docx import Document
import httpx

import main
from fallback import CAB_ST, CAB_CPP, MILEAGE_ST, MILEAGE_CPP, fallback_code
from file_import import extract_document
from llm_client import LLMClient, generation_messages
from scenarios import SCENARIOS
from validation import CHECKS, local_report, structural_issues, merge_review, repair_needed, improved_report


class FakeModel:
    configured = True
    broken = False
    stream_calls = 0
    review_calls = 0

    async def stream(self, messages):
        self.stream_calls += 1
        if self.broken:
            raise RuntimeError("upstream_error")
        code = CAB_CPP if "C++17" in messages[0]["content"] else CAB_ST
        for start in range(0, len(code), 120):
            yield code[start:start + 120]
            await asyncio.sleep(0)

    async def complete(self, messages):
        return CAB_ST

    async def review(self, requirement, code, language):
        self.review_calls += 1
        if self.broken:
            raise RuntimeError("review_error")
        return {"checks": [{"id": key, "status": "pass", "detail": "测试替身返回，非真实模型结论"} for key, _ in CHECKS], "coverage": [{"requirement": "测试项", "status": "pass", "evidence": "测试替身"}], "suggestions": []}

    async def close(self):
        pass


def events(response):
    return [json.loads(line[5:].strip()) for line in response.text.splitlines() if line.startswith("data:")]


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.fake = FakeModel()
        self.patcher = patch.object(main, "llm", self.fake)
        self.patcher.start()
        main.cache.clear()
        main.rate_log.clear()
        self.client = TestClient(main.app)

    def tearDown(self):
        self.client.close()
        self.patcher.stop()

    def test_health_no_credentials(self):
        data = self.client.get("/health").json()
        self.assertEqual(data["version"], "9.0.0")
        self.assertFalse(any(key in data for key in ["api_key", "model", "api_base"]))

    def test_frontend_served_same_origin(self):
        result = self.client.get("/")
        self.assertEqual(result.status_code, 200)
        self.assertIn("代码工作台", result.text)
        self.assertNotIn("Access-Control-Allow-Origin", result.headers)
        self.assertEqual(result.headers["x-frame-options"], "DENY")
        self.assertEqual(self.client.get("/static/app.js").status_code, 200)

    def test_six_scenarios(self):
        data = self.client.get("/api/scenarios").json()["scenarios"]
        self.assertEqual(len(data), 6)
        self.assertEqual(data[-1]["title"], "轮缘润滑及测试")

    def test_stream_and_partial_before_final(self):
        result = events(self.client.post("/api/generate-stream", json={"requirement": "操作端判断", "language": "st"}))
        types = [r["type"] for r in result]
        self.assertIn("delta", types)
        self.assertLess(types.index("checks"), types.index("code"))
        self.assertLess(types.index("code"), types.index("done"))
        partial = next(e["report"] for e in result if e["type"] == "checks")
        self.assertEqual(partial["source"], "partial")
        self.assertTrue(all(c["status"] == "pending" for c in partial["checks"]))
        final = result[-1]["report"]
        self.assertEqual(len(final["checks"]), 6)
        self.assertEqual({c["id"] for c in final["checks"]}, {i for i, _ in CHECKS})

    def test_cpp_route_and_cache_isolation(self):
        for lang in ("st", "cpp"):
            r = self.client.post("/api/generate-code", json={"requirement": "操作端", "language": lang}).json()
            self.assertEqual(r["language"], lang)
            self.assertIn("FUNCTION_BLOCK" if lang == "st" else "class CabSelection", r["code"])
        self.assertEqual(self.fake.stream_calls, 2)
        r = self.client.post("/api/generate-code", json={"requirement": "操作端", "language": "cpp"}).json()
        self.assertEqual(r["source"], "cache")
        self.assertEqual(self.fake.stream_calls, 2)
        self.assertNotEqual(main.cache_key("变量 xA", "st"), main.cache_key("变量 xa", "st"))

    def test_model_failure_no_irrelevant_template(self):
        self.fake.broken = True
        result = events(self.client.post("/api/generate-stream", json={"requirement": "PID温控 37摄氏度", "language": "st"}))
        self.assertEqual(result[-1]["type"], "error")
        self.assertFalse(any(e["type"] == "code" for e in result))
        self.assertEqual(len(main.cache), 0)

    def test_exact_reference_has_source_and_no_cache(self):
        self.fake.broken = True
        r = self.client.post("/api/generate-code", json={"requirement": SCENARIOS[0]["requirement"], "language": "cpp"}).json()
        self.assertEqual(r["source"], "reference")
        self.assertEqual(r["report"]["source"], "rules")
        self.assertGreater(r["report"]["risk_count"], 0)
        self.assertEqual(len(main.cache), 0)

    def test_greeting_is_not_code(self):
        r = self.client.post("/api/generate-code", json={"requirement": "你好"}).json()
        self.assertEqual(r["mode"], "chat")
        self.assertFalse(r["can_validate"])
        self.assertEqual(self.fake.stream_calls, 0)

    def test_validation_uses_same_six_items(self):
        data = self.client.post("/api/validate-code", json={"requirement": "操作端", "code": CAB_ST, "language": "st"}).json()
        self.assertEqual([c["id"] for c in data["checks"]], [k for k, _ in CHECKS])

    def setup_improvement(self, mode="better"):
        fake = self.fake
        fake.improve_calls = 0
        async def review(requirement, code, language):
            fake.review_calls += 1
            checks = [{"id": key, "status": "pass", "detail": "有对应实现", "repairable": False} for key, _ in CHECKS]
            if fake.review_calls == 1:
                checks[2].update(status="warn", detail="声明需修正", repairable=mode != "ambiguous")
            if mode == "regression" and fake.review_calls > 1:
                checks[0].update(status="fail", detail="新的语法问题")
            return {"checks": checks, "coverage": [{"requirement": "操作端", "evidence": "状态判断", "status": "pass"}], "suggestions": []}
        async def improve(requirement, code, language, report):
            fake.improve_calls += 1
            if mode == "unavailable":
                raise RuntimeError("test_only")
            if mode == "invalid":
                return "FUNCTION_BLOCK Incomplete"
            return code + "\n(* reviewed candidate *)"
        fake.review = review
        fake.improve = improve

    def test_improvement_is_reviewed_before_acceptance_and_cached(self):
        self.setup_improvement()
        data = events(self.client.post("/api/generate-stream", json={"requirement": "操作端", "language": "st"}))
        codes = [e for e in data if e["type"] == "code"]
        self.assertEqual(len(codes), 2)
        self.assertNotEqual(codes[0]["code_id"], codes[1]["code_id"])
        self.assertEqual(data[-1]["code_id"], codes[1]["code_id"])
        self.assertTrue(data[-1]["report"]["quality_improvement"]["accepted"])
        self.assertEqual(main.cache[main.cache_key("操作端", "st")][1], codes[-1]["code"])
        self.assertEqual(self.fake.improve_calls, 1)
        self.assertEqual(self.fake.review_calls, 2)

    def test_failed_or_regressing_repair_keeps_original_findings(self):
        for mode in ("unavailable", "invalid", "regression"):
            with self.subTest(mode=mode):
                main.cache.clear(); self.fake.review_calls = 0
                self.setup_improvement(mode)
                data = events(self.client.post("/api/generate-stream", json={"requirement": "操作端", "language": "st"}))
                self.assertEqual(sum(e["type"] == "code" for e in data), 1)
                self.assertFalse(data[-1]["report"]["quality_improvement"]["accepted"])
                self.assertEqual(data[-1]["report"]["checks"][2]["status"], "warn")

    def test_ambiguous_requirement_not_silently_rewritten(self):
        self.setup_improvement("ambiguous")
        result = self.client.post("/api/generate-code", json={"requirement": "操作端"}).json()
        self.assertEqual(self.fake.improve_calls, 0)
        self.assertEqual(result["report"]["checks"][2]["status"], "warn")

    def test_validate_only_never_changes_code(self):
        self.setup_improvement()
        result = events(self.client.post("/api/validate-stream", json={"requirement": "操作端", "language": "st", "code": CAB_ST}))
        self.assertEqual(self.fake.improve_calls, 0)
        self.assertFalse(any(e["type"] == "code" for e in result))

    def test_compact_v9_frontend(self):
        html = self.client.get("/").text
        self.assertIn("代码验证", html)
        self.assertIn("workspace.css?v=9", html)
        self.assertNotIn("看得见的验证过程", html)
        self.assertNotIn("让控制需求，成为可审查的代码", html)
        self.assertIn('id="reviewDisclosure"', html)

    def test_bad_inputs(self):
        for body in [{"requirement": " "}, {"requirement": "x", "language": "python"}, {"requirement": "x" * 20001}]:
            self.assertEqual(self.client.post("/api/generate-code", json=body).status_code, 422)

    def test_text_upload(self):
        response = self.client.post("/api/import-file", files={"file": ("需求.txt", "采样周期 0.1 秒".encode(), "text/plain")})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["text"], "采样周期 0.1 秒")

    def test_docx_upload_order(self):
        doc = Document()
        doc.add_paragraph("前置条件")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "参数 5 秒"
        doc.add_paragraph("最后停止")
        b = io.BytesIO(); doc.save(b)
        response = self.client.post("/api/import-file", files={"file": ("需求.docx", b.getvalue())})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["text"], "前置条件\n参数 5 秒\n最后停止")

    def test_pdf_text_upload(self):
        from pypdf import PdfWriter
        from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 10 270 Td (Input period 0.1 seconds) Tj ET")
        page[NameObject("/Contents")] = content
        stream = io.BytesIO(); writer.write(stream)
        response = self.client.post("/api/import-file", files={"file": ("requirement.pdf", stream.getvalue())})
        self.assertEqual(response.status_code, 200)
        self.assertIn("Input period 0.1 seconds", response.json()["text"])

    def test_pdf_without_text_is_not_silent_success(self):
        from pypdf import PdfWriter
        writer = PdfWriter(); writer.add_blank_page(width=300, height=300)
        stream = io.BytesIO(); writer.write(stream)
        self.assertEqual(self.client.post("/api/import-file", files={"file": ("scan.pdf", stream.getvalue())}).status_code, 400)

    def test_upload_limits(self):
        for filename, data in [("需求.doc", b"data"), ("坏.txt", b"\0\0"), ("大.txt", b"x" * (2 * 1024 * 1024 + 1)), ("空.txt", b""), ("长.txt", b"x" * 20001)]:
            self.assertEqual(self.client.post("/api/import-file", files={"file": (filename, data)}).status_code, 400)

    def test_oversized_request_rejected_before_parsing(self):
        self.assertEqual(self.client.post("/api/import-file", content=b"x" * (3 * 1024 * 1024)).status_code, 413)


class RulesTests(unittest.TestCase):
    def test_missing_review_never_becomes_all_pass(self):
        report = merge_review(local_report(CAB_ST, "st"), {})
        self.assertEqual(report["source"], "rules")
        self.assertGreater(report["risk_count"], 0)

    def test_local_failure_retained_despite_short_presentation(self):
        report = merge_review(local_report(CAB_ST.replace("END_FUNCTION_BLOCK", ""), "st"), {"checks": [{"id": k, "status": "pass", "detail": "有依据", "repairable": False} for k, _ in CHECKS]})
        self.assertTrue(repair_needed(report))
        self.assertEqual(report["checks"][0]["status"], "fail")

    def test_missing_coverage_is_not_a_pass(self):
        report = merge_review(local_report(CAB_ST, "st"), {"checks": [{"id": k, "status": "pass", "detail": "有依据"} for k, _ in CHECKS]})
        self.assertEqual(report["checks"][4]["status"], "warn")

    def test_references_complete(self):
        for lang, code in [("st", CAB_ST), ("st", MILEAGE_ST), ("cpp", CAB_CPP), ("cpp", MILEAGE_CPP)]:
            self.assertEqual(structural_issues(code, lang), [])

    def test_crossed_st_blocks_fail(self):
        code = CAB_ST.replace("END_IF;\nCASE", "END_CASE;\nCASE", 1)
        self.assertTrue(structural_issues(code, "st"))

    def test_missing_end_fails(self):
        self.assertTrue(structural_issues(CAB_ST.replace("END_FUNCTION_BLOCK", ""), "st"))

    def test_st_keywords_inside_string_not_blocks(self):
        code = CAB_ST.replace("iState : INT := 0;", "iState : INT := 0;\nsMessage : STRING := 'IF CASE END_VAR';")
        self.assertEqual(structural_issues(code, "st"), [])

    def test_cpp_unbalanced_fails(self):
        self.assertTrue(structural_issues(CAB_CPP + "}", "cpp"))

    def test_local_failure_cannot_be_overruled(self):
        local = local_report(CAB_ST.replace("END_FUNCTION_BLOCK", ""), "st")
        report = merge_review(local, {"checks": [{"id": k, "status": "pass", "detail": "模型认为通过"} for k, _ in CHECKS]})
        self.assertEqual(report["checks"][0]["status"], "fail")

    def test_modified_requirement_never_gets_reference(self):
        self.assertIsNone(fallback_code(SCENARIOS[0]["requirement"] + "\n增加停车检测", "st"))

    def test_prompts_preserve_requirement_and_language(self):
        messages = generation_messages("每 7 秒停止；不使用默认 5 秒", "cpp")
        self.assertIn("C++17", messages[0]["content"])
        self.assertIn("每 7 秒停止", messages[1]["content"])

    def test_coverage_failure_matches_check_summary(self):
        report = merge_review(local_report(CAB_ST, "st"), {"checks": [{"id": "requirements", "status": "pass", "detail": "无问题"}], "coverage": [{"requirement": "保护条件", "status": "fail", "evidence": "没有实现"}]})
        self.assertEqual(next(x for x in report["checks"] if x["id"] == "requirements")["status"], "fail")


class ProviderTests(unittest.IsolatedAsyncioTestCase):
    async def test_stream_parameter_retry_and_truncation(self):
        calls = []
        async def transport(request):
            body = json.loads(request.content); calls.append(body)
            if "thinking" in body:
                return httpx.Response(400, json={"message": "unsupported"})
            return httpx.Response(200, headers={"content-type": "text/event-stream"}, text='data: {"choices":[{"delta":{"content":"partial"},"finish_reason":null}]}\n\ndata: {"choices":[{"delta":{},"finish_reason":"length"}]}\n\n')
        model = LLMClient("https://model.invalid/v1", "test-only", "demo")
        model.client = httpx.AsyncClient(transport=httpx.MockTransport(transport))
        with self.assertRaises(Exception):
            async for _ in model.stream(generation_messages("测试", "st")):
                pass
        self.assertEqual(len(calls), 2)
        self.assertNotIn("thinking", calls[1])
        await model.close()

    async def test_stream_disconnect_not_complete(self):
        async def transport(request):
            return httpx.Response(200, text='data: {"choices":[{"delta":{"content":"partial"}}]}\n\n')
        model = LLMClient("https://model.invalid/v1", "test-only", "demo", client=httpx.AsyncClient(transport=httpx.MockTransport(transport)))
        with self.assertRaises(Exception):
            async for _ in model.stream(generation_messages("测试", "st")):
                pass
        await model.close()


if __name__ == "__main__":
    unittest.main()
